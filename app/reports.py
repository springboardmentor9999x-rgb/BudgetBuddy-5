from datetime import datetime, date
from decimal import Decimal
from io import BytesIO

from fastapi import (
    APIRouter,
    Depends,
    HTTPException,
    Query,
)
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.database import get_db
from app.models import (
    Report,
    User,
    BankAccount,
    Income,
    Expense,
    Budget,
    SavingsGoal,
    SavingsTransaction,
)
from app.users import get_current_user


router = APIRouter(
    prefix="/reports",
    tags=["Reports"]
)


# ==========================================================
# VALID MODULES
# ==========================================================

VALID_MODULES = {
    "financial_summary",
    "income_transactions",
    "expense_transactions",
    "income_categories",
    "expense_categories",
    "budget_report",
    "savings_goals",
    "savings_transactions",
    "bank_account_details",
}


MODULE_NAMES = {
    "financial_summary": "Financial Summary",
    "income_transactions": "Income Transactions",
    "expense_transactions": "Expense Transactions",
    "income_categories": "Income Categories",
    "expense_categories": "Expense Categories",
    "budget_report": "Budget Report",
    "savings_goals": "Savings Goals",
    "savings_transactions": "Savings Transactions",
    "bank_account_details": "Bank Account Details",
}


# ==========================================================
# HELPERS
# ==========================================================

def money(value):
    if value is None:
        return 0.0

    if isinstance(value, Decimal):
        return float(value)

    return float(value)


def format_money(value):
    return f"₹{money(value):,.2f}"


def format_date(value):
    if value is None:
        return "-"

    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y")

    return value.strftime("%d-%m-%Y")


def format_datetime(value):
    if value is None:
        return "-"

    return value.strftime("%d-%m-%Y %I:%M:%S %p")


def format_time(value):
    if value is None:
        return "-"

    return value.strftime("%I:%M:%S %p")


def get_selected_modules(modules: str):

    selected_modules = [
        item.strip()
        for item in modules.split(",")
        if item.strip()
    ]

    # Remove duplicates while preserving order
    selected_modules = list(
        dict.fromkeys(selected_modules)
    )

    invalid_modules = [
        item
        for item in selected_modules
        if item not in VALID_MODULES
    ]

    if invalid_modules:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid modules: {invalid_modules}"
        )

    if not selected_modules:
        raise HTTPException(
            status_code=400,
            detail="Please select at least one report module"
        )

    return selected_modules


# ==========================================================
# GET SELECTED BANK
# ==========================================================

def get_selected_bank(
    bank_account_id,
    db,
    current_user
):

    if bank_account_id is None:
        return None

    bank = (
        db.query(BankAccount)
        .filter(
            BankAccount.id == bank_account_id,
            BankAccount.user_id == current_user.id
        )
        .first()
    )

    if bank is None:
        raise HTTPException(
            status_code=404,
            detail="Bank account not found"
        )

    return bank


# ==========================================================
# FINANCIAL SUMMARY DATA
# ==========================================================

def build_financial_summary(
    db,
    current_user,
    bank_account_id
):

    income_query = (
        db.query(
            func.coalesce(
                func.sum(Income.amount),
                0
            )
        )
        .filter(
            Income.user_id == current_user.id
        )
    )

    expense_query = (
        db.query(
            func.coalesce(
                func.sum(Expense.amount),
                0
            )
        )
        .filter(
            Expense.user_id == current_user.id
        )
    )

    savings_query = (
        db.query(
            func.coalesce(
                func.sum(SavingsTransaction.amount),
                0
            )
        )
        .filter(
            SavingsTransaction.user_id
            == current_user.id
        )
    )

    if bank_account_id is not None:

        income_query = income_query.filter(
            Income.bank_account_id
            == bank_account_id
        )

        expense_query = expense_query.filter(
            Expense.bank_account_id
            == bank_account_id
        )

        savings_query = savings_query.filter(
            SavingsTransaction.bank_account_id
            == bank_account_id
        )

    total_income = money(
        income_query.scalar()
    )

    total_expenses = money(
        expense_query.scalar()
    )

    total_saved = money(
        savings_query.scalar()
    )

    balance = (
        total_income
        - total_expenses
        - total_saved
    )

    return {
        "Total Income": total_income,
        "Total Expenses": total_expenses,
        "Total Amount Saved": total_saved,
        "Balance": balance,
    }


# ==========================================================
# PREVIEW REPORT
# ==========================================================

@router.get("/preview")
def preview_report(

    bank_account_id: int | None = None,

    modules: str = Query(
        ...,
        description=(
            "Comma separated report modules"
        )
    ),

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    )
):

    selected_modules = get_selected_modules(
        modules
    )

    selected_bank = get_selected_bank(
        bank_account_id,
        db,
        current_user
    )

    bank_name = (
        selected_bank.bank_name
        if selected_bank
        else "All Bank Accounts"
    )

    account_number = (
        selected_bank.account_number
        if selected_bank
        else None
    )

    result = {
        "report_for": bank_name,
        "account_number": account_number,
        "generated_on": format_datetime(
            datetime.now()
        ),
        "modules": []
    }


    # ======================================================
    # FINANCIAL SUMMARY
    # ======================================================

    if "financial_summary" in selected_modules:

        summary = build_financial_summary(
            db,
            current_user,
            bank_account_id
        )

        result["financial_summary"] = summary

        result["modules"].append(
            "Financial Summary"
        )


    # ======================================================
    # INCOME TRANSACTIONS
    # ======================================================

    if "income_transactions" in selected_modules:

        query = (
            db.query(Income)
            .filter(
                Income.user_id
                == current_user.id
            )
            .order_by(
                Income.created_at.desc()
            )
        )

        if bank_account_id is not None:

            query = query.filter(
                Income.bank_account_id
                == bank_account_id
            )

        incomes = query.all()

        rows = []

        total = 0

        for income in incomes:

            amount = money(
                income.amount
            )

            total += amount

            rows.append({
                "Date": format_date(
                    income.date
                ),
                "Exact Time": format_time(
                    income.created_at
                ),
                "Source": income.source,
                "Category": income.category,
                "Amount": amount,
                "Description": (
                    income.description
                    or ""
                )
            })

        result["income_transactions"] = {
            "rows": rows,
            "total": total
        }

        result["modules"].append(
            "Income Transactions"
        )


    # ======================================================
    # EXPENSE TRANSACTIONS
    # ======================================================

    if "expense_transactions" in selected_modules:

        query = (
            db.query(Expense)
            .filter(
                Expense.user_id
                == current_user.id
            )
            .order_by(
                Expense.created_at.desc()
            )
        )

        if bank_account_id is not None:

            query = query.filter(
                Expense.bank_account_id
                == bank_account_id
            )

        expenses = query.all()

        rows = []

        total = 0

        for expense in expenses:

            amount = money(
                expense.amount
            )

            total += amount

            rows.append({
                "Date": format_date(
                    expense.date
                ),
                "Exact Time": format_time(
                    expense.created_at
                ),
                "Category": expense.category,
                "Payment Method": (
                    expense.payment_method
                ),
                "Amount": amount,
                "Description": (
                    expense.description
                    or ""
                )
            })

        result["expense_transactions"] = {
            "rows": rows,
            "total": total
        }

        result["modules"].append(
            "Expense Transactions"
        )


    # ======================================================
    # INCOME CATEGORIES
    # ======================================================

    if "income_categories" in selected_modules:

        query = (
            db.query(
                Income.category,
                func.sum(
                    Income.amount
                )
            )
            .filter(
                Income.user_id
                == current_user.id
            )
        )

        if bank_account_id is not None:

            query = query.filter(
                Income.bank_account_id
                == bank_account_id
            )

        category_data = (
            query
            .group_by(
                Income.category
            )
            .order_by(
                func.sum(
                    Income.amount
                ).desc()
            )
            .all()
        )

        rows = []

        total = 0

        for category, amount in category_data:

            amount = money(amount)

            total += amount

            rows.append({
                "Category": category,
                "Total Income": amount
            })

        result["income_categories"] = {
            "rows": rows,
            "total": total
        }

        result["modules"].append(
            "Income Categories"
        )


    # ======================================================
    # EXPENSE CATEGORIES
    # ======================================================

    if "expense_categories" in selected_modules:

        query = (
            db.query(
                Expense.category,
                func.sum(
                    Expense.amount
                )
            )
            .filter(
                Expense.user_id
                == current_user.id
            )
        )

        if bank_account_id is not None:

            query = query.filter(
                Expense.bank_account_id
                == bank_account_id
            )

        category_data = (
            query
            .group_by(
                Expense.category
            )
            .order_by(
                func.sum(
                    Expense.amount
                ).desc()
            )
            .all()
        )

        rows = []

        total = 0

        for category, amount in category_data:

            amount = money(amount)

            total += amount

            rows.append({
                "Category": category,
                "Total Spent": amount
            })

        result["expense_categories"] = {
            "rows": rows,
            "total": total
        }

        result["modules"].append(
            "Expense Categories"
        )


    # ======================================================
    # BUDGET REPORT
    # ======================================================

    if "budget_report" in selected_modules:

        query = (
            db.query(Budget)
            .filter(
                Budget.user_id
                == current_user.id
            )
            .order_by(
                Budget.year.desc(),
                Budget.month.desc()
            )
        )

        if bank_account_id is not None:

            query = query.filter(
                Budget.bank_account_id
                == bank_account_id
            )

        budgets = query.all()

        rows = []

        total_spent = 0

        for budget in budgets:

            start_date = date(
                budget.year,
                budget.month,
                1
            )

            if budget.month == 12:

                end_date = date(
                    budget.year + 1,
                    1,
                    1
                )

            else:

                end_date = date(
                    budget.year,
                    budget.month + 1,
                    1
                )

            expense_query = (
                db.query(
                    func.coalesce(
                        func.sum(
                            Expense.amount
                        ),
                        0
                    )
                )
                .filter(
                    Expense.user_id
                    == current_user.id,

                    func.lower(
                        func.trim(
                            Expense.category
                        )
                    )
                    ==
                    budget.category.strip().lower(),

                    Expense.date >= start_date,

                    Expense.date < end_date
                )
            )

            if bank_account_id is not None:

                expense_query = (
                    expense_query.filter(
                        Expense.bank_account_id
                        == bank_account_id
                    )
                )

            spent = money(
                expense_query.scalar()
            )

            limit = money(
                budget.monthly_limit
            )

            remaining = limit - spent

            percentage = (
                (spent / limit) * 100
                if limit > 0
                else 0
            )

            total_spent += spent

            rows.append({
                "Category": budget.category,
                "Month": budget.month,
                "Year": budget.year,
                "Limit": limit,
                "Spent": spent,
                "Remaining": remaining,
                "Percentage Used": round(
                    percentage,
                    2
                )
            })

        result["budget_report"] = {
            "rows": rows,
            "total_spent": total_spent
        }

        result["modules"].append(
            "Budget Report"
        )


    # ======================================================
    # SAVINGS GOALS
    # ======================================================

    if "savings_goals" in selected_modules:

        query = (
            db.query(SavingsGoal)
            .filter(
                SavingsGoal.user_id
                == current_user.id
            )
            .order_by(
                SavingsGoal.created_at.desc()
            )
        )

        if bank_account_id is not None:

            query = query.filter(
                SavingsGoal.bank_account_id
                == bank_account_id
            )

        goals = query.all()

        rows = []

        total_target = 0
        total_saved = 0

        for goal in goals:

            target = money(
                goal.target_amount
            )

            saved = money(
                goal.current_amount
            )

            remaining = max(
                target - saved,
                0
            )

            progress = (
                (saved / target) * 100
                if target > 0
                else 0
            )

            total_target += target
            total_saved += saved

            rows.append({
                "Goal": goal.goal_name,
                "Target": target,
                "Saved": saved,
                "Remaining": remaining,
                "Progress": round(
                    progress,
                    2
                )
            })

        result["savings_goals"] = {
            "rows": rows,
            "total_target": total_target,
            "total_saved": total_saved,
            "remaining": max(
                total_target - total_saved,
                0
            )
        }

        result["modules"].append(
            "Savings Goals"
        )


    # ======================================================
    # SAVINGS TRANSACTIONS
    # ======================================================

    if "savings_transactions" in selected_modules:

        query = (
            db.query(
                SavingsTransaction,
                SavingsGoal
            )
            .join(
                SavingsGoal,
                SavingsTransaction.goal_id
                == SavingsGoal.id
            )
            .filter(
                SavingsTransaction.user_id
                == current_user.id
            )
            .order_by(
                SavingsTransaction.created_at.desc()
            )
        )

        if bank_account_id is not None:

            query = query.filter(
                SavingsTransaction.bank_account_id
                == bank_account_id
            )

        transactions = query.all()

        rows = []

        total = 0

        for transaction, goal in transactions:

            amount = money(
                transaction.amount
            )

            total += amount

            rows.append({
                "Date": format_date(
                    transaction.transaction_date
                ),
                "Exact Time": format_time(
                    transaction.created_at
                ),
                "Goal": goal.goal_name,
                "Amount": amount
            })

        result["savings_transactions"] = {
            "rows": rows,
            "total": total
        }

        result["modules"].append(
            "Savings Transactions"
        )


    # ======================================================
    # BANK ACCOUNT DETAILS
    # ======================================================

    if "bank_account_details" in selected_modules:

        if selected_bank:

            accounts = [
                selected_bank
            ]

        else:

            accounts = (
                db.query(BankAccount)
                .filter(
                    BankAccount.user_id
                    == current_user.id
                )
                .order_by(
                    BankAccount.bank_name
                )
                .all()
            )

        rows = []

        for account in accounts:

            rows.append({
                "Bank": account.bank_name,
                "Account Holder":
                    account.account_holder,
                "Account Number":
                    account.account_number,
                "IFSC":
                    account.ifsc_code,
                "Type":
                    account.account_type,
                "Balance":
                    money(
                        account.current_balance
                    )
            })

        result["bank_account_details"] = {
            "rows": rows
        }

        result["modules"].append(
            "Bank Account Details"
        )


    return result


# ==========================================================
# SAVE REPORT HISTORY
# ==========================================================

def save_report_history(
    db,
    current_user,
    selected_bank,
    selected_modules
):

    bank_text = (
        selected_bank.bank_name
        if selected_bank
        else "All Bank Accounts"
    )

    module_text = ", ".join(
        MODULE_NAMES.get(
            module,
            module
        )
        for module in selected_modules
    )

    # Keep report_type safely below VARCHAR(50)
    report_type = (
        "Financial Report"
    )

    report = Report(
        user_id=current_user.id,
        bank_account_id=(
            selected_bank.id
            if selected_bank
            else None
        ),
        report_type=report_type,
        generated_date=datetime.now()
    )

    db.add(report)
    db.commit()
    db.refresh(report)

    return report


# ==========================================================
# CREATE REPORT HISTORY
# ==========================================================

@router.post("")
def create_report(
    report_type: str = Query(
        "Financial Report"
    ),

    bank_account_id: int | None = None,

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    )
):

    selected_bank = get_selected_bank(
        bank_account_id,
        db,
        current_user
    )

    report = Report(
        user_id=current_user.id,
        bank_account_id=(
            selected_bank.id
            if selected_bank
            else None
        ),
        report_type="Financial Report",
        generated_date=datetime.now()
    )

    db.add(report)
    db.commit()
    db.refresh(report)

    return report


# ==========================================================
# GET REPORT HISTORY
# ==========================================================

@router.get("")
def get_reports(
    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    )
):

    reports = (
        db.query(Report)
        .filter(
            Report.user_id
            == current_user.id
        )
        .order_by(
            Report.generated_date.desc()
        )
        .all()
    )

    return reports


# ==========================================================
# BUILD PDF
# ==========================================================

def build_pdf(
    report_data,
    selected_modules
):

    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=20,
        spaceAfter=15,
    )

    section_style = ParagraphStyle(
        "SectionTitle",
        parent=styles["Heading2"],
        fontSize=15,
        spaceBefore=12,
        spaceAfter=8,
    )

    story = []

    story.append(
        Paragraph(
            "BudgetBuddy Financial Report",
            title_style
        )
    )

    story.append(
        Paragraph(
            f"<b>Report For:</b> "
            f"{report_data['report_for']}"
        )
    )

    if report_data.get(
        "account_number"
    ):

        story.append(
            Paragraph(
                f"<b>Account Number:</b> "
                f"{report_data['account_number']}"
            )
        )

    story.append(
        Paragraph(
            f"<b>Generated On:</b> "
            f"{report_data['generated_on']}"
        )
    )

    story.append(
        Spacer(1, 15)
    )


    def add_table(
        data,
        widths=None
    ):

        table = Table(
            data,
            colWidths=widths,
            repeatRows=1
        )

        table.setStyle(
            TableStyle([
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor(
                        "#2563eb"
                    )
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE"
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor(
                            "#f8fafc"
                        )
                    ]
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6
                ),
            ])
        )

        story.append(table)
        story.append(
            Spacer(1, 10)
        )


    # ======================================================
    # FINANCIAL SUMMARY
    # ======================================================

    if "financial_summary" in selected_modules:

        story.append(
            Paragraph(
                "Financial Summary",
                section_style
            )
        )

        summary = (
            report_data[
                "financial_summary"
            ]
        )

        data = [
            [
                "Financial Summary",
                "Amount"
            ],
            [
                "Total Income",
                format_money(
                    summary["Total Income"]
                )
            ],
            [
                "Total Expenses",
                format_money(
                    summary["Total Expenses"]
                )
            ],
            [
                "Total Amount Saved",
                format_money(
                    summary[
                        "Total Amount Saved"
                    ]
                )
            ],
            [
                "Balance",
                format_money(
                    summary["Balance"]
                )
            ],
        ]

        add_table(
            data,
            [
                90 * mm,
                70 * mm
            ]
        )


    # ======================================================
    # INCOME TRANSACTIONS
    # ======================================================

    if "income_transactions" in selected_modules:

        story.append(
            Paragraph(
                "Income Transactions",
                section_style
            )
        )

        info = report_data[
            "income_transactions"
        ]

        data = [
            [
                "Date",
                "Exact Time",
                "Source",
                "Category",
                "Amount"
            ]
        ]

        for row in info["rows"]:

            data.append([
                row["Date"],
                row["Exact Time"],
                row["Source"],
                row["Category"],
                format_money(
                    row["Amount"]
                )
            ])

        if len(data) == 1:

            data.append([
                "No transactions",
                "",
                "",
                "",
                ""
            ])

        add_table(
            data,
            [
                25 * mm,
                28 * mm,
                30 * mm,
                32 * mm,
                30 * mm
            ]
        )

        story.append(
            Paragraph(
                f"<b>Total Income:</b> "
                f"{format_money(info['total'])}"
            )
        )


    # ======================================================
    # EXPENSE TRANSACTIONS
    # ======================================================

    if "expense_transactions" in selected_modules:

        story.append(
            Paragraph(
                "Expense Transactions",
                section_style
            )
        )

        info = report_data[
            "expense_transactions"
        ]

        data = [
            [
                "Date",
                "Exact Time",
                "Category",
                "Payment Method",
                "Amount"
            ]
        ]

        for row in info["rows"]:

            data.append([
                row["Date"],
                row["Exact Time"],
                row["Category"],
                row["Payment Method"],
                format_money(
                    row["Amount"]
                )
            ])

        if len(data) == 1:

            data.append([
                "No transactions",
                "",
                "",
                "",
                ""
            ])

        add_table(
            data,
            [
                25 * mm,
                28 * mm,
                35 * mm,
                40 * mm,
                30 * mm
            ]
        )

        story.append(
            Paragraph(
                f"<b>Total Amount Spent:</b> "
                f"{format_money(info['total'])}"
            )
        )


    # ======================================================
    # INCOME CATEGORIES
    # ======================================================

    if "income_categories" in selected_modules:

        story.append(
            Paragraph(
                "Income Categories",
                section_style
            )
        )

        info = report_data[
            "income_categories"
        ]

        data = [
            [
                "Category",
                "Total Income"
            ]
        ]

        for row in info["rows"]:

            data.append([
                row["Category"],
                format_money(
                    row["Total Income"]
                )
            ])

        if len(data) == 1:

            data.append([
                "No income categories",
                ""
            ])

        add_table(
            data,
            [
                90 * mm,
                70 * mm
            ]
        )

        story.append(
            Paragraph(
                f"<b>Total Income:</b> "
                f"{format_money(info['total'])}"
            )
        )


    # ======================================================
    # EXPENSE CATEGORIES
    # ======================================================

    if "expense_categories" in selected_modules:

        story.append(
            Paragraph(
                "Expense Categories",
                section_style
            )
        )

        info = report_data[
            "expense_categories"
        ]

        data = [
            [
                "Category",
                "Total Spent"
            ]
        ]

        for row in info["rows"]:

            data.append([
                row["Category"],
                format_money(
                    row["Total Spent"]
                )
            ])

        if len(data) == 1:

            data.append([
                "No expense categories",
                ""
            ])

        add_table(
            data,
            [
                90 * mm,
                70 * mm
            ]
        )

        story.append(
            Paragraph(
                f"<b>Total Amount Spent:</b> "
                f"{format_money(info['total'])}"
            )
        )


    # ======================================================
    # BUDGET REPORT
    # ======================================================

    if "budget_report" in selected_modules:

        story.append(
            Paragraph(
                "Budget Report",
                section_style
            )
        )

        info = report_data[
            "budget_report"
        ]

        data = [
            [
                "Category",
                "Month",
                "Year",
                "Limit",
                "Spent",
                "Remaining"
            ]
        ]

        for row in info["rows"]:

            data.append([
                row["Category"],
                str(row["Month"]),
                str(row["Year"]),
                format_money(
                    row["Limit"]
                ),
                format_money(
                    row["Spent"]
                ),
                format_money(
                    row["Remaining"]
                )
            ])

        if len(data) == 1:

            data.append([
                "No budgets",
                "",
                "",
                "",
                "",
                ""
            ])

        add_table(
            data,
            [
                35 * mm,
                20 * mm,
                18 * mm,
                28 * mm,
                28 * mm,
                30 * mm
            ]
        )

        story.append(
            Paragraph(
                f"<b>Total Amount Spent:</b> "
                f"{format_money(info['total_spent'])}"
            )
        )


    # ======================================================
    # SAVINGS GOALS
    # ======================================================

    if "savings_goals" in selected_modules:

        story.append(
            Paragraph(
                "Savings Goals",
                section_style
            )
        )

        info = report_data[
            "savings_goals"
        ]

        data = [
            [
                "Goal",
                "Target",
                "Saved",
                "Remaining",
                "Progress"
            ]
        ]

        for row in info["rows"]:

            data.append([
                row["Goal"],
                format_money(
                    row["Target"]
                ),
                format_money(
                    row["Saved"]
                ),
                format_money(
                    row["Remaining"]
                ),
                f"{row['Progress']:.2f}%"
            ])

        if len(data) == 1:

            data.append([
                "No savings goals",
                "",
                "",
                "",
                ""
            ])

        add_table(
            data,
            [
                45 * mm,
                30 * mm,
                30 * mm,
                30 * mm,
                25 * mm
            ]
        )

        story.append(
            Paragraph(
                f"<b>Total Savings Target:</b> "
                f"{format_money(info['total_target'])}"
            )
        )

        story.append(
            Paragraph(
                f"<b>Total Amount Saved:</b> "
                f"{format_money(info['total_saved'])}"
            )
        )

        story.append(
            Paragraph(
                f"<b>Savings Remaining:</b> "
                f"{format_money(info['remaining'])}"
            )
        )


    # ======================================================
    # SAVINGS TRANSACTIONS
    # ======================================================

    if "savings_transactions" in selected_modules:

        story.append(
            Paragraph(
                "Savings Transactions",
                section_style
            )
        )

        info = report_data[
            "savings_transactions"
        ]

        data = [
            [
                "Date",
                "Exact Time",
                "Goal",
                "Amount"
            ]
        ]

        for row in info["rows"]:

            data.append([
                row["Date"],
                row["Exact Time"],
                row["Goal"],
                format_money(
                    row["Amount"]
                )
            ])

        if len(data) == 1:

            data.append([
                "No savings transactions",
                "",
                "",
                ""
            ])

        add_table(
            data,
            [
                35 * mm,
                35 * mm,
                55 * mm,
                35 * mm
            ]
        )

        story.append(
            Paragraph(
                f"<b>Total Amount Saved:</b> "
                f"{format_money(info['total'])}"
            )
        )


    # ======================================================
    # BANK ACCOUNT DETAILS
    # ======================================================

    if "bank_account_details" in selected_modules:

        story.append(
            Paragraph(
                "Bank Account Details",
                section_style
            )
        )

        info = report_data[
            "bank_account_details"
        ]

        data = [
            [
                "Bank",
                "Account Holder",
                "Account Number",
                "IFSC",
                "Type",
                "Balance"
            ]
        ]

        for row in info["rows"]:

            data.append([
                row["Bank"],
                row["Account Holder"],
                row["Account Number"],
                row["IFSC"],
                row["Type"],
                format_money(
                    row["Balance"]
                )
            ])

        if len(data) == 1:

            data.append([
                "No bank accounts",
                "",
                "",
                "",
                "",
                ""
            ])

        add_table(
            data,
            [
                30 * mm,
                32 * mm,
                35 * mm,
                28 * mm,
                25 * mm,
                30 * mm
            ]
        )


    document.build(story)

    buffer.seek(0)

    return buffer


# ==========================================================
# DOWNLOAD PDF
# ==========================================================

@router.get("/download/pdf")
def download_pdf(

    bank_account_id: int | None = None,

    modules: str = Query(...),

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    )
):

    selected_modules = get_selected_modules(
        modules
    )

    selected_bank = get_selected_bank(
        bank_account_id,
        db,
        current_user
    )

    report_data = preview_report(
        bank_account_id=bank_account_id,
        modules=modules,
        db=db,
        current_user=current_user
    )

    buffer = build_pdf(
        report_data,
        selected_modules
    )

    save_report_history(
        db,
        current_user,
        selected_bank,
        selected_modules
    )

    filename = (
        "BudgetBuddy_Financial_Report.pdf"
    )

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"'
        }
    )


# ==========================================================
# EXCEL HELPERS
# ==========================================================

def add_excel_sheet(
    workbook,
    sheet_name,
    headers,
    rows,
    total_label=None,
    total_value=None
):

    # Excel sheet names have a 31 character limit
    sheet = workbook.create_sheet(
        title=sheet_name[:31]
    )

    # Header
    for col_index, header in enumerate(
        headers,
        start=1
    ):

        cell = sheet.cell(
            row=1,
            column=col_index,
            value=header
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="2563EB"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # Rows
    for row_index, row in enumerate(
        rows,
        start=2
    ):

        for col_index, value in enumerate(
            row,
            start=1
        ):

            sheet.cell(
                row=row_index,
                column=col_index,
                value=value
            )

    # Total
    if total_label is not None:

        row_number = (
            len(rows) + 2
        )

        sheet.cell(
            row=row_number,
            column=1,
            value=total_label
        ).font = Font(
            bold=True
        )

        sheet.cell(
            row=row_number,
            column=2,
            value=total_value
        ).font = Font(
            bold=True
        )

    # Formatting
    for column_cells in sheet.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            try:

                length = len(
                    str(cell.value)
                )

                max_length = max(
                    max_length,
                    length
                )

            except Exception:
                pass

        sheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 3,
            40
        )

    sheet.freeze_panes = "A2"

    return sheet


# ==========================================================
# BUILD EXCEL
# ==========================================================

def build_excel(
    report_data,
    selected_modules
):

    workbook = Workbook()

    # Remove default sheet
    default_sheet = workbook.active

    workbook.remove(
        default_sheet
    )


    # ======================================================
    # REPORT INFORMATION
    # ======================================================

    info_sheet = workbook.create_sheet(
        "Report Info"
    )

    info_sheet["A1"] = (
        "BudgetBuddy Financial Report"
    )

    info_sheet["A1"].font = Font(
        bold=True,
        size=16
    )

    info_sheet["A3"] = "Report For"

    info_sheet["B3"] = (
        report_data["report_for"]
    )

    info_sheet["A4"] = (
        "Account Number"
    )

    info_sheet["B4"] = (
        report_data.get(
            "account_number"
        )
        or "All Accounts"
    )

    info_sheet["A5"] = (
        "Generated On"
    )

    info_sheet["B5"] = (
        report_data["generated_on"]
    )

    info_sheet.column_dimensions[
        "A"
    ].width = 25

    info_sheet.column_dimensions[
        "B"
    ].width = 40


    # ======================================================
    # FINANCIAL SUMMARY
    # ======================================================

    if "financial_summary" in selected_modules:

        summary = report_data[
            "financial_summary"
        ]

        rows = [
            [
                "Total Income",
                summary["Total Income"]
            ],
            [
                "Total Expenses",
                summary["Total Expenses"]
            ],
            [
                "Total Amount Saved",
                summary[
                    "Total Amount Saved"
                ]
            ],
            [
                "Balance",
                summary["Balance"]
            ]
        ]

        add_excel_sheet(
            workbook,
            "Financial Summary",
            [
                "Description",
                "Amount"
            ],
            rows
        )


    # ======================================================
    # INCOME TRANSACTIONS
    # ======================================================

    if "income_transactions" in selected_modules:

        info = report_data[
            "income_transactions"
        ]

        rows = []

        for row in info["rows"]:

            rows.append([
                row["Date"],
                row["Exact Time"],
                row["Source"],
                row["Category"],
                row["Amount"],
                row["Description"]
            ])

        add_excel_sheet(
            workbook,
            "Income Transactions",
            [
                "Date",
                "Exact Time",
                "Source",
                "Category",
                "Amount",
                "Description"
            ],
            rows,
            "Total Income",
            info["total"]
        )


    # ======================================================
    # EXPENSE TRANSACTIONS
    # ======================================================

    if "expense_transactions" in selected_modules:

        info = report_data[
            "expense_transactions"
        ]

        rows = []

        for row in info["rows"]:

            rows.append([
                row["Date"],
                row["Exact Time"],
                row["Category"],
                row["Payment Method"],
                row["Amount"],
                row["Description"]
            ])

        add_excel_sheet(
            workbook,
            "Expense Transactions",
            [
                "Date",
                "Exact Time",
                "Category",
                "Payment Method",
                "Amount",
                "Description"
            ],
            rows,
            "Total Amount Spent",
            info["total"]
        )


    # ======================================================
    # INCOME CATEGORIES
    # ======================================================

    if "income_categories" in selected_modules:

        info = report_data[
            "income_categories"
        ]

        rows = []

        for row in info["rows"]:

            rows.append([
                row["Category"],
                row["Total Income"]
            ])

        add_excel_sheet(
            workbook,
            "Income Categories",
            [
                "Category",
                "Total Income"
            ],
            rows,
            "Total Income",
            info["total"]
        )


    # ======================================================
    # EXPENSE CATEGORIES
    # ======================================================

    if "expense_categories" in selected_modules:

        info = report_data[
            "expense_categories"
        ]

        rows = []

        for row in info["rows"]:

            rows.append([
                row["Category"],
                row["Total Spent"]
            ])

        add_excel_sheet(
            workbook,
            "Expense Categories",
            [
                "Category",
                "Total Spent"
            ],
            rows,
            "Total Amount Spent",
            info["total"]
        )


    # ======================================================
    # BUDGET REPORT
    # ======================================================

    if "budget_report" in selected_modules:

        info = report_data[
            "budget_report"
        ]

        rows = []

        for row in info["rows"]:

            rows.append([
                row["Category"],
                row["Month"],
                row["Year"],
                row["Limit"],
                row["Spent"],
                row["Remaining"],
                row["Percentage Used"]
            ])

        add_excel_sheet(
            workbook,
            "Budget Report",
            [
                "Category",
                "Month",
                "Year",
                "Limit",
                "Spent",
                "Remaining",
                "Percentage Used"
            ],
            rows,
            "Total Amount Spent",
            info["total_spent"]
        )


    # ======================================================
    # SAVINGS GOALS
    # ======================================================

    if "savings_goals" in selected_modules:

        info = report_data[
            "savings_goals"
        ]

        rows = []

        for row in info["rows"]:

            rows.append([
                row["Goal"],
                row["Target"],
                row["Saved"],
                row["Remaining"],
                row["Progress"]
            ])

        add_excel_sheet(
            workbook,
            "Savings Goals",
            [
                "Goal",
                "Target",
                "Saved",
                "Remaining",
                "Progress %"
            ],
            rows
        )

        sheet = workbook[
            "Savings Goals"
        ]

        next_row = len(rows) + 3

        sheet.cell(
            next_row,
            1,
            "Total Savings Target"
        ).font = Font(
            bold=True
        )

        sheet.cell(
            next_row,
            2,
            info["total_target"]
        ).font = Font(
            bold=True
        )

        sheet.cell(
            next_row + 1,
            1,
            "Total Amount Saved"
        ).font = Font(
            bold=True
        )

        sheet.cell(
            next_row + 1,
            2,
            info["total_saved"]
        ).font = Font(
            bold=True
        )

        sheet.cell(
            next_row + 2,
            1,
            "Savings Remaining"
        ).font = Font(
            bold=True
        )

        sheet.cell(
            next_row + 2,
            2,
            info["remaining"]
        ).font = Font(
            bold=True
        )


    # ======================================================
    # SAVINGS TRANSACTIONS
    # ======================================================

    if "savings_transactions" in selected_modules:

        info = report_data[
            "savings_transactions"
        ]

        rows = []

        for row in info["rows"]:

            rows.append([
                row["Date"],
                row["Exact Time"],
                row["Goal"],
                row["Amount"]
            ])

        add_excel_sheet(
            workbook,
            "Savings Transactions",
            [
                "Date",
                "Exact Time",
                "Goal",
                "Amount"
            ],
            rows,
            "Total Amount Saved",
            info["total"]
        )


    # ======================================================
    # BANK ACCOUNT DETAILS
    # ======================================================

    if "bank_account_details" in selected_modules:

        info = report_data[
            "bank_account_details"
        ]

        rows = []

        for row in info["rows"]:

            rows.append([
                row["Bank"],
                row["Account Holder"],
                row["Account Number"],
                row["IFSC"],
                row["Type"],
                row["Balance"]
            ])

        add_excel_sheet(
            workbook,
            "Bank Account Details",
            [
                "Bank",
                "Account Holder",
                "Account Number",
                "IFSC",
                "Type",
                "Balance"
            ],
            rows
        )


    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output


# ==========================================================
# DOWNLOAD EXCEL
# ==========================================================

@router.get("/download/excel")
def download_excel(

    bank_account_id: int | None = None,

    modules: str = Query(...),

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    )
):

    selected_modules = get_selected_modules(
        modules
    )

    selected_bank = get_selected_bank(
        bank_account_id,
        db,
        current_user
    )

    report_data = preview_report(
        bank_account_id=bank_account_id,
        modules=modules,
        db=db,
        current_user=current_user
    )

    buffer = build_excel(
        report_data,
        selected_modules
    )

    save_report_history(
        db,
        current_user,
        selected_bank,
        selected_modules
    )

    filename = (
        "BudgetBuddy_Financial_Report.xlsx"
    )

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"'
        }
    )


# ==========================================================
# GET ONE REPORT
# ==========================================================

@router.get("/{report_id}")
def get_report(
    report_id: int,

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    )
):

    report = (
        db.query(Report)
        .filter(
            Report.id == report_id,
            Report.user_id
            == current_user.id
        )
        .first()
    )

    if report is None:

        raise HTTPException(
            status_code=404,
            detail="Report not found"
        )

    return report


# ==========================================================
# DELETE REPORT
# ==========================================================

@router.delete("/{report_id}")
def delete_report(
    report_id: int,

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    )
):

    report = (
        db.query(Report)
        .filter(
            Report.id == report_id,
            Report.user_id
            == current_user.id
        )
        .first()
    )

    if report is None:

        raise HTTPException(
            status_code=404,
            detail="Report not found"
        )

    db.delete(report)
    db.commit()

    return {
        "message":
            "Report deleted successfully"
    }